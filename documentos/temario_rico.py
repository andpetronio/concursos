"""Texto enriquecido acotado para temarios; no guarda ni ejecuta HTML arbitrario.

Solo admite párrafos, viñetas, numeración, negrita, cursiva y subrayado.
Se admite texto plano histórico como formato de entrada.
"""
from html.parser import HTMLParser
import re
from copy import deepcopy

from django.core.exceptions import ValidationError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor

PERMITIDOS = {'p', 'div', 'br', 'ul', 'ol', 'li', 'b', 'strong', 'i', 'em', 'u'}


class _Lector(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.bloques = []
        self.actual = []
        self.tipo = 'texto'
        self.lista = []
        self.enlinea = []
        self.tags = []
        self.total = 0

    def _cerrar(self):
        if any(texto.strip() for texto, *_ in self.actual):
            self.bloques.append((self.tipo, self.actual))
        self.actual = []
        self.tipo = 'texto'

    def handle_starttag(self, tag, attrs):
        if tag not in PERMITIDOS or attrs:
            raise ValidationError('El temario contiene formato no admitido. Usá negrita, cursiva, subrayado o listas.')
        if tag in ('ul', 'ol'):
            self._cerrar()
            if self.lista:
                raise ValidationError('El temario admite listas de un solo nivel.')
            self.lista.append('vineta' if tag == 'ul' else 'numero')
        elif tag in ('p', 'div', 'li'):
            self._cerrar()
            self.tipo = self.lista[-1] if tag == 'li' and self.lista else 'texto'
        elif tag == 'br':
            self.actual.append(('\n', False, False, False))
        else:
            self.enlinea.append(tag)
        if tag != 'br':
            self.tags.append(tag)

    def handle_endtag(self, tag):
        if tag not in PERMITIDOS or tag == 'br':
            raise ValidationError('El formato del temario no es válido.')
        if not self.tags or self.tags[-1] != tag:
            raise ValidationError('El formato del temario tiene etiquetas incompletas.')
        self.tags.pop()
        if tag in ('p', 'div', 'li'):
            self._cerrar()
        elif tag in ('ul', 'ol'):
            self._cerrar()
            self.lista.pop()
        else:
            self.enlinea.pop()

    def handle_data(self, data):
        if not data:
            return
        self.total += len(data)
        self.actual.append((data, bool({'strong', 'b'} & set(self.enlinea)),
                            bool({'em', 'i'} & set(self.enlinea)), 'u' in self.enlinea))

    def finish(self):
        if self.tags or self.lista:
            raise ValidationError('El formato del temario quedó incompleto.')
        self._cerrar()
        if self.total > 4000:
            raise ValidationError('El temario no puede superar 4000 caracteres de contenido.')
        if not self.bloques:
            raise ValidationError('Ingresá el contenido del temario sorteado.')
        return self.bloques


def leer_temario(valor):
    valor = str(valor or '').strip()
    if not valor:
        raise ValidationError('Ingresá el contenido del temario sorteado.')
    if len(valor) > 16000:
        raise ValidationError('El temario es demasiado extenso.')
    if not re.search(r'<[a-zA-Z/][^>]*>', valor):
        bloques = [('texto', [(linea, False, False, False)]) for linea in valor.splitlines() if linea.strip()]
        if len(valor) > 4000 or not bloques:
            raise ValidationError('El temario no puede superar 4000 caracteres de contenido.')
        return bloques
    lector = _Lector()
    try:
        lector.feed(valor)
        lector.close()
        return lector.finish()
    except (ValueError, TypeError) as exc:
        raise ValidationError('El formato del temario no es válido.') from exc


def _numeracion_decimal(documento):
    root = documento.part.numbering_part.element
    actuales = [int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))]
    abstract_id = max(actuales, default=-1) + 1
    nums = [int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))]
    num_id = max(nums, default=0) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
    for nombre, valor in (('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')):
        node = OxmlElement('w:' + nombre); node.set(qn('w:val'), valor); lvl.append(node)
    ppr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '720'); ind.set(qn('w:hanging'), '360'); ppr.append(ind); lvl.append(ppr)
    abstract.append(lvl)
    # abstractNum se ubica antes de los w:num para respetar el esquema OOXML.
    primero_num = root.find(qn('w:num'))
    if primero_num is not None: primero_num.addprevious(abstract)
    else: root.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(abstract_id))
    num.append(ref); root.append(num)
    return num_id


def _poner_numero(parrafo, num_id):
    ppr = parrafo._p.get_or_add_pPr()
    existente = ppr.find(qn('w:numPr'))
    if existente is not None:
        ppr.remove(existente)
    if num_id is not None:
        numpr = OxmlElement('w:numPr')
        nivel = OxmlElement('w:ilvl'); nivel.set(qn('w:val'), '0')
        numero = OxmlElement('w:numId'); numero.set(qn('w:val'), str(num_id))
        numpr.append(nivel); numpr.append(numero)
        ppr.insert(0, numpr)


def poner_temario(documento, indice, valor, num_vineta):
    """Sustituye una viñeta de ejemplo por N párrafos Word editables y formateados."""
    bloques = leer_temario(valor)
    original = documento.paragraphs[indice]
    referencia = original
    numeracion = None
    for posicion, (tipo, fragmentos) in enumerate(bloques):
        if posicion == 0:
            parrafo = original
        else:
            elemento = deepcopy(original._p)
            referencia._p.addnext(elemento)
            parrafo = Paragraph(elemento, original._parent)
        parrafo.clear()  # Conserva sangría y espacio del modelo, no los textos de ejemplo.
        if tipo == 'numero':
            if numeracion is None: numeracion = _numeracion_decimal(documento)
            _poner_numero(parrafo, numeracion)
        else:
            _poner_numero(parrafo, num_vineta if tipo == 'vineta' else None)
            if tipo != 'numero': numeracion = None
        for texto, negrita, cursiva, subrayado in fragmentos:
            partes = texto.split('\n')
            for i, parte in enumerate(partes):
                if i: parrafo.add_run().add_break()
                if parte:
                    run = parrafo.add_run(parte)
                    run.bold = negrita; run.italic = cursiva; run.underline = subrayado
                    run.font.color.rgb = RGBColor(0, 0, 0)
        referencia = parrafo
