"""Generación documental por falta de aspirantes: requisitos y datos de expediente."""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document

from concursos.models import Concurso
from concursos.services import crear_concurso
from concursos.tests.test_concursos import DATOS, PASSWORD
from participantes.models import Aspirante, Persona
from . import services
from .forms import GeneracionForm
from .models import GeneracionDocumento
from .registry import POR_CODIGO, TIPOS


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class DesiertoSinInscriptosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.operador = get_user_model().objects.create_user("operador_desierto", password=PASSWORD)
        cls.operador.user_permissions.set(Permission.objects.filter(
            content_type__app_label__in=("concursos", "documentos")))
        cls.lector = get_user_model().objects.create_user("lector_desierto", password=PASSWORD)
        cls.lector.user_permissions.add(Permission.objects.get(
            content_type__app_label="concursos", codename="view_concurso"))

    def setUp(self):
        self.concurso = crear_concurso(
            datos={**DATOS, "jurisdiccion": "rectorado"}, usuario=self.operador)
        self.concurso.resolucion_convocatoria = "3278/26-R"
        self.concurso.save(update_fields=["resolucion_convocatoria"])
        self.tipo = POR_CODIGO["desierto_inscriptos"]
        # El login real inicializa SessionActivity, que force_login no crea.
        acceso = self.client.post(reverse("login"), {
            "username": self.operador.username, "password": PASSWORD,
        })
        self.assertEqual(acceso.status_code, 302)

    def confirmar_acta(self):
        acta = self.concurso.controles_documentales.get(requisito="acta_inscriptos")
        acta.fojas = "14/15"
        acta.verificado = True
        acta.save(update_fields=["fojas", "verificado"])

    def generar(self, *, huella=None, confirmar=True, usuario=None):
        actual = Concurso.objects.get(pk=self.concurso.pk)
        snapshot, _, _, _ = services.fuentes_actuales(actual, self.tipo)
        return services.generar(
            concurso_id=actual.pk, tipo_codigo=self.tipo.codigo,
            usuario=usuario or self.operador, version_concurso=actual.version,
            version_destinatario=0, huella_esperada=huella or services.huella_fuentes(snapshot),
            complementarios={"confirmacion_acta_sin_inscriptos": "on"} if confirmar else {},
        )

    def test_ficha_genera_directamente_sin_vista_de_revision(self):
        self.confirmar_acta()
        url = reverse("documentos:generar_en_ficha", args=[self.concurso.pk, self.tipo.codigo])
        ficha = self.client.get(self.concurso.get_absolute_url())
        self.assertEqual(ficha.status_code, 200)
        self.assertContains(ficha, f'action="{url}#documentos"')
        estado = next(d for d in ficha.context["tipos_documentos"] if d["tipo"] == self.tipo)
        formulario = estado["formulario"]
        # Un input oculto sin valor se envía como cadena vacía desde el navegador.
        # Preservar también el cero de los campos de control de versiones.
        datos = {
            field.html_name: "" if field.value() is None else field.value()
            for field in formulario.hidden_fields()
        }
        datos[formulario.add_prefix("confirmacion_acta_sin_inscriptos")] = "on"
        respuesta = self.client.post(url, datos)
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(GeneracionDocumento.objects.count(), 1)
        self.assertIn("/documentos/archivo/", respuesta["Location"])

    def test_ficha_muestra_error_de_confirmacion_sin_perder_datos(self):
        self.confirmar_acta()
        url = reverse("documentos:generar_en_ficha", args=[self.concurso.pk, self.tipo.codigo])
        ficha = self.client.get(self.concurso.get_absolute_url())
        estado = next(d for d in ficha.context["tipos_documentos"] if d["tipo"] == self.tipo)
        formulario = estado["formulario"]
        # Un input oculto sin valor se envía como cadena vacía desde el navegador.
        # Preservar también el cero de los campos de control de versiones.
        datos = {
            field.html_name: "" if field.value() is None else field.value()
            for field in formulario.hidden_fields()
        }
        respuesta = self.client.post(url, datos)
        self.assertEqual(respuesta.status_code, 400)
        self.assertContains(respuesta, "panel-documentos", status_code=400)
        estado = next(d for d in respuesta.context["tipos_documentos"] if d["tipo"] == self.tipo)
        self.assertIn("confirmacion_acta_sin_inscriptos", estado["formulario"].errors)
        self.assertFalse(GeneracionDocumento.objects.exists())

    def test_lector_no_puede_generar_por_url_directa(self):
        self.confirmar_acta()
        self.client.logout()
        self.client.post(reverse("login"), {"username": self.lector.username, "password": PASSWORD})
        respuesta = self.client.post(reverse("documentos:generar_en_ficha", args=[
            self.concurso.pk, self.tipo.codigo]), {})
        self.assertEqual(respuesta.status_code, 403)
        self.assertFalse(GeneracionDocumento.objects.exists())

    def test_catalogo_ocho_documentos_pero_sin_nueva_migracion(self):
        self.assertEqual(len(TIPOS), 8)
        self.assertEqual(self.tipo.version, "desierto-inscriptos-1")
        self.assertEqual(self.tipo.fuente, "")

    def test_cero_cargados_sin_acta_no_acredita_concurso_desierto(self):
        estado = services.disponibilidad(self.concurso, self.tipo)
        self.assertFalse(estado["disponible"])
        self.assertTrue(any("acta" in valor for valor in estado["faltantes"]))
        with self.assertRaisesMessage(ValidationError, "acta"):
            self.generar()
        self.assertFalse(GeneracionDocumento.objects.exists())

    def test_exige_jurisdiccion_para_generar_sin_inventar_datos(self):
        self.confirmar_acta()
        self.concurso.jurisdiccion = ""
        self.concurso.save(update_fields=["jurisdiccion"])
        estado = services.disponibilidad(self.concurso, self.tipo)
        self.assertFalse(estado["disponible"])
        self.assertTrue(any("jurisdicción" in x for x in estado["faltantes"]))
        with self.assertRaisesMessage(ValidationError, "jurisdicción"):
            self.generar()
        self.assertFalse(GeneracionDocumento.objects.exists())

    def test_no_habilita_cuando_hay_inscriptos_aunque_acta_este_verificada(self):
        self.confirmar_acta()
        persona = Persona.objects.create(apellido="De prueba", nombre="Aspirante",
            dni="12345678", creado_por=self.operador, actualizado_por=self.operador)
        Aspirante.objects.create(concurso=self.concurso, persona=persona,
            creado_por=self.operador, actualizado_por=self.operador)
        estado = services.disponibilidad(self.concurso, self.tipo)
        self.assertTrue(any("Hay aspirantes" in x for x in estado["faltantes"]))
        with self.assertRaisesMessage(ValidationError, "Hay aspirantes"):
            self.generar()

    def test_exige_resolucion_y_modalidad_con_nuevo_llamado_definido(self):
        self.confirmar_acta()
        for modalidad in ("abierto", "acotado"):
            self.concurso.modalidad = modalidad
            self.concurso.save(update_fields=["modalidad"])
            self.assertFalse(services.disponibilidad(self.concurso, self.tipo)["disponible"])
        self.concurso.modalidad = "cerrado_interno"
        self.concurso.resolucion_convocatoria = ""
        self.concurso.save(update_fields=["modalidad", "resolucion_convocatoria"])
        self.assertFalse(services.disponibilidad(self.concurso, self.tipo)["disponible"])
        with self.assertRaisesMessage(ValidationError, "resolución oficial"):
            self.generar()

    def test_generacion_exige_confirmacion_acta_y_preserva_estado(self):
        self.confirmar_acta()
        self.assertTrue(services.disponibilidad(self.concurso, self.tipo)["disponible"])
        formulario = GeneracionForm(tipo=self.tipo, concurso=self.concurso)
        self.assertIn("confirmacion_acta_sin_inscriptos", formulario.fields)
        with self.assertRaisesMessage(ValidationError, "confirmacion_acta_sin_inscriptos"):
            self.generar(confirmar=False)
        registro = self.generar()
        documento = Document(BytesIO(bytes(registro.contenido)))
        texto = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("3278/26-R", texto)
        self.assertIn(self.concurso.expediente, texto)
        self.assertIn("14/15", texto)
        self.assertIn("Cerrado General", texto)
        self.assertIn("ARTÍCULO 1°.-", texto)
        self.assertIn("ARTÍCULO 2°.-", texto)
        self.assertNotIn("PUNTAJE", texto)
        self.assertEqual(registro.instantanea["nuevo_llamado"]["modalidad"], "cerrado_general")
        self.assertEqual(registro.plantilla_version, "desierto-inscriptos-1")
        self.concurso.refresh_from_db()
        self.assertEqual(self.concurso.estado, "preparacion")
        self.assertEqual(Concurso.objects.count(), 1)

    def test_cerrado_general_genera_borrador_que_sugiere_abierto(self):
        self.confirmar_acta()
        self.concurso.modalidad = "cerrado_general"
        self.concurso.save(update_fields=["modalidad"])
        registro = self.generar()
        self.assertEqual(registro.instantanea["nuevo_llamado"]["nombre"], "Abierto")
        self.assertIn("Abierto", "\n".join(p.text for p in
            Document(BytesIO(bytes(registro.contenido))).paragraphs))

    def test_fuentes_cambiadas_y_permisos_impiden_generacion(self):
        self.confirmar_acta()
        snapshot, _, _, _ = services.fuentes_actuales(self.concurso, self.tipo)
        huella = services.huella_fuentes(snapshot)
        self.concurso.resolucion_convocatoria = "9999/26-R"
        self.concurso.save(update_fields=["resolucion_convocatoria"])
        with self.assertRaisesMessage(ValidationError, "Los datos cambiaron"):
            self.generar(huella=huella)
        with self.assertRaises(PermissionDenied):
            self.generar(usuario=self.lector)

    def test_revision_web_no_se_confunde_con_resolucion_emitida(self):
        self.confirmar_acta()
        url = reverse("documentos:revisar", args=[self.concurso.pk, self.tipo.codigo])
        pagina = self.client.get(url)
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, "Confirmo que el acta")
        self.assertContains(pagina, "no crea ni convoca otro concurso")
        self.client.logout()
        acceso = self.client.post(reverse("login"), {
            "username": self.lector.username, "password": PASSWORD,
        })
        self.assertEqual(acceso.status_code, 302)
        self.assertEqual(self.client.get(url).status_code, 403)
