"""Regresiones del reemplazo de variables en modelos Word institucionales."""

from django.test import SimpleTestCase
from django.core.exceptions import ValidationError
from pathlib import Path
from docx import Document

from .renderers_modelos import _reemplazar, _reemplazar_literal_obligatorio, _designacion


class ReemplazoModelosWordTests(SimpleTestCase):
    def test_ignora_sustitucion_identica_sin_bucle(self):
        documento = Document()
        parrafo = documento.add_paragraph('Tramo Inicial, Agrupamiento Administrativo 01.')
        _reemplazar(parrafo, {
            'Tramo Inicial': 'Tramo Inicial',
            'Administrativo 01': 'Administrativo 01',
        })
        self.assertEqual(parrafo.text, 'Tramo Inicial, Agrupamiento Administrativo 01.')

    def test_reemplazo_que_contiene_el_original_se_procesa_una_vez(self):
        documento = Document()
        parrafo = documento.add_paragraph('Resolución 3164/23-R.')
        _reemplazar(parrafo, {'3164/23-R': 'Antecedente 3164/23-R'})
        self.assertEqual(parrafo.text, 'Resolución Antecedente 3164/23-R.')

    def test_sustituye_todas_las_ocurrencias_incluso_entre_runs(self):
        documento = Document()
        parrafo = documento.add_paragraph()
        parrafo.add_run('El {{nom')
        parrafo.add_run('bre}} y {{nombre}}').bold = True
        _reemplazar(parrafo, {'{{nombre}}': 'Pérez'})
        self.assertEqual(parrafo.text, 'El Pérez y Pérez')


class IntegridadDeModelosWordTests(SimpleTestCase):
    """Una actualización del modelo no puede dejar datos de ejemplo sin detectar."""

    def test_literal_obligatorio_se_reemplaza_aunque_word_lo_divida_en_runs(self):
        documento = Document()
        parrafo = documento.add_paragraph()
        parrafo.add_run('Hora: doce')
        parrafo.add_run(' horas.')
        _reemplazar_literal_obligatorio(parrafo, 'doce horas', '11:40 horas', 'cierre')
        self.assertEqual(parrafo.text, 'Hora: 11:40 horas.')

    def test_literal_obligatorio_ausente_bloquea_antes_de_generar(self):
        documento = Document()
        parrafo = documento.add_paragraph('El cierre se registra a las 11:40 horas.')
        with self.assertRaisesMessage(ValidationError, 'Cambió el modelo institucional (hora de cierre)'):
            _reemplazar_literal_obligatorio(parrafo, 'doce horas', '12:00 horas', 'hora de cierre')

    @staticmethod
    def _datos_designacion():
        complementarios = {
            'aspirante_id': '1', 'aplica_limitacion': True,
            'resolucion_anterior': '1234/26-R', 'categoria_anterior': '6',
            'tramo_anterior': 'Intermedio', 'agrupamiento_anterior': 'Servicios 03',
            'cp_anterior': '10888', 'funciones_anteriores': 'Archivo',
            'unidad_anterior': 'Secretaría Administrativa',
        }
        for campo in ('ejercicio', 'grupo', 'ppal', 'subppal', 'subsubppal', 'fte',
                      'prog', 'sprog', 'proy', 'activ', 'obra', 'finalidad', 'funcion',
                      'unidad', 'unidad_cargo', 'planta', 'inciso', 'partida',
                      'caracter', 'escalafon'):
            complementarios['presupuesto_' + campo] = '01'
        return {
            'complementarios': complementarios,
            'concurso': {'causa_vacancia': 'jubilacion', 'fecha_baja': '2026-10-20',
                         'ex_titular': 'Ex Titular', 'resolucion_baja': '1234/26-R',
                         'acto_vacancia': '', 'detalle_vacancia': ''},
            'aspirantes': [{'id': 1, 'persona': {
                'nombre': 'Ana', 'apellido': 'Ejemplo', 'dni': '12345678'}}],
        }

    @staticmethod
    def _modelo_designacion():
        return Document(Path(__file__).with_name('modelos') / 'PROYECTO RES. DESIGNACION.docx')

    def test_designacion_usa_causa_real_y_no_hereda_jubilacion_del_modelo(self):
        escenarios = {
            'nuevo': 'nueva creación',
            'traslado': 'como consecuencia de su traslado',
            'jubilacion': 'beneficio jubilatorio',
            'renuncia': 'presentó su renuncia',
            'fallecimiento': 'fallecimiento',
            'concurso': 'en otro cargo por concurso',
            'efectivizacion': 'efectivización',
            'otra': 'Otra situación acreditada en el expediente',
        }
        for causa, esperado in escenarios.items():
            with self.subTest(causa=causa):
                documento = self._modelo_designacion()
                datos = self._datos_designacion()
                datos['concurso']['causa_vacancia'] = causa
                datos['concurso']['acto_vacancia'] = 'Resolución 5678/26-R'
                datos['concurso']['detalle_vacancia'] = 'Otra situación acreditada en el expediente'
                _designacion(documento, datos, {})
                considerandos = [p.text for p in documento.paragraphs
                                 if p.text.startswith('Que el cargo')]
                self.assertEqual(len(considerandos), 1)
                self.assertIn(esperado, considerandos[0])
                if causa != 'jubilacion':
                    self.assertNotIn('beneficio jubilatorio', considerandos[0])
                    self.assertNotIn('su renuncia definitiva fue aceptada', considerandos[0])

    def test_designacion_rechaza_plantilla_con_considerando_de_vacancia_modificado(self):
        documento = self._modelo_designacion()
        for parrafo in documento.paragraphs:
            if parrafo.text.startswith('Que según lo informado en las referidas actuaciones'):
                _reemplazar(parrafo, {'beneficio jubilatorio': 'otra causal de ejemplo'})
                break
        with self.assertRaisesMessage(ValidationError, 'Cambió el considerando de vacancia'):
            _designacion(documento, self._datos_designacion(), {})

    def test_designacion_sustituye_los_datos_de_ejemplo_de_la_limitacion(self):
        documento = self._modelo_designacion()
        m = {}
        _designacion(documento, self._datos_designacion(), m)
        bloques = [p.text for p in documento.paragraphs if 'Limitar' in p.text or 'corresponde limitar' in p.text]
        self.assertEqual(len(bloques), 2)
        for texto in bloques:
            self.assertIn('1234/26-R', texto)
            self.assertNotIn('3164/23-R', texto)
            self.assertNotIn('11991', texto)
        self.assertEqual(m['{{ganador_a}}'], 'Ana Ejemplo')

    def test_designacion_no_acepta_modelo_con_dato_de_ejemplo_modificado(self):
        documento = self._modelo_designacion()
        for parrafo in documento.paragraphs:
            if parrafo.text.startswith('Que corresponde limitar,'):
                _reemplazar(parrafo, {'Departamento de Rendición de Cuentas': 'Departamento de ejemplo'})
                break
        with self.assertRaisesMessage(ValidationError, 'dato de ejemplo'):
            _designacion(documento, self._datos_designacion(), {})

    def test_designacion_sin_limitacion_quita_ambos_bloques(self):
        documento = self._modelo_designacion()
        datos = self._datos_designacion()
        datos['complementarios']['aplica_limitacion'] = False
        _designacion(documento, datos, {})
        self.assertFalse(any(p.text.startswith('Que corresponde limitar,') or
                             p.text.startswith('ARTÍCULO 2º.- Limitar') for p in documento.paragraphs))
        self.assertTrue(any(p.text.startswith('ARTÍCULO 2º.- Cumplido') for p in documento.paragraphs))
