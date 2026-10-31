"""Interpolación fiel de los tres modelos Word provistos por la Dirección de Concursos.

Los originales son fuentes inmutables. Solo se sustituyen variables y secciones
repetibles/condicionales explícitas, sin recrear ni reescribir los considerandos.
"""
from io import BytesIO
from pathlib import Path
import re

from django.conf import settings
from django.core.exceptions import ValidationError
from docx import Document

from concursos.catalogos import CATALOGO_CARGOS
from concursos.models import Concurso
from .renderers import causa_de_vacancia
from .temario_rico import poner_temario

FUENTES_WORD = {
    "sorteo_acta": "Acta de sorteo de temas.docx",
    "dictamen": "PROYECTO RES. APROBANDO DICTAMEN.docx",
    "designacion": "PROYECTO RES. DESIGNACION.docx",
}
PATRON = re.compile(r"\{\{[^{}]*\}\}")


def _nombre(registro):
    p = registro["persona"]
    return f'{p["nombre"]} {p["apellido"]}'.strip()


def _dni(registro):
    dni = str(registro["persona"]["dni"])
    return f'{dni[:2]}.{dni[2:5]}.{dni[5:]}' if len(dni) == 8 and dni.isdecimal() else dni


def _fecha(valor):
    # El formulario DateField entrega ISO a la instantánea histórica.
    if not valor:
        return ""
    from datetime import date
    d = date.fromisoformat(str(valor)[:10])
    return f'{d.day} de {("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")[d.month - 1]} de {d.year}'


def _fecha_acta(valor):
    """Primera línea del acta: «29 días del mes de octubre de 2026»."""
    fecha = _fecha(valor)
    if not fecha:
        return ''
    dia, resto = fecha.split(' de ', 1)
    return f'{dia} {"día" if dia == "1" else "días"} del mes de {resto}'


def _reemplazar(parrafo, cambios):
    """Reemplaza marcadores incluso cuando Word los fragmentó entre varios runs."""
    for viejo, nuevo in cambios.items():
        reemplazo = str(nuevo)
        if not viejo or viejo == reemplazo:
            continue
        # Avanzar después del texto insertado: podría contener el patrón original.
        # Volver a buscar desde cero dejaría este bucle ejecutándose indefinidamente.
        desde = 0
        while True:
            runs = list(parrafo.runs)
            completo = ''.join(run.text for run in runs)
            idx = completo.find(viejo, desde)
            if idx < 0:
                break
            fin = idx + len(viejo)
            pos = 0
            comienzo = cierre = None
            start = end = 0
            for i, run in enumerate(runs):
                prox = pos + len(run.text)
                if comienzo is None and pos <= idx < prox:
                    comienzo, start = i, idx - pos
                if cierre is None and pos < fin <= prox:
                    cierre, end = i, fin - pos
                pos = prox
            if comienzo is None or cierre is None:
                raise ValidationError('Un marcador Word no pudo sustituirse de forma íntegra.')
            antes = runs[comienzo].text[:start]
            despues = runs[cierre].text[end:]
            if comienzo == cierre:
                runs[comienzo].text = antes + reemplazo + despues
            else:
                runs[comienzo].text = antes + reemplazo
                for i in range(comienzo + 1, cierre):
                    runs[i].text = ''
                runs[cierre].text = despues
            desde = idx + len(reemplazo)


def _reemplazar_literal_obligatorio(parrafo, original, nuevo, descripcion):
    """Si cambia un literal de ejemplo del modelo Word, nunca emitirlo sin sustituir."""
    if original not in parrafo.text:
        raise ValidationError(
            f'Cambió el modelo institucional ({descripcion}); revisá el archivo fuente antes de generar.'
        )
    _reemplazar(parrafo, {original: nuevo})


def _quitar(parrafo):
    parrafo._element.getparent().remove(parrafo._element)


def _lista_personas(registros, con_dni=False):
    nombres = [(_nombre(a) + (f', {_dni(a)}' if con_dni else '')) for a in registros]
    if not nombres:
        return ''
    if len(nombres) == 1:
        return nombres[0]
    return ', '.join(nombres[:-1]) + ' y ' + nombres[-1]


def _cantidad_letras(n):
    numeros = ('CERO', 'UN', 'DOS', 'TRES', 'CUATRO', 'CINCO', 'SEIS', 'SIETE', 'OCHO', 'NUEVE',
               'DIEZ', 'ONCE', 'DOCE', 'TRECE', 'CATORCE', 'QUINCE', 'DIECISÉIS', 'DIECISIETE',
               'DIECIOCHO', 'DIECINUEVE', 'VEINTE')
    return numeros[n] if n <= 20 else str(n)


def _jurado(j, tratamiento=''):
    return ' '.join(v for v in (tratamiento.strip(), j['titulo'].strip(), _nombre(j)) if v)


def _comunes(datos):
    c, x = datos['concurso'], datos['complementarios']
    agrupamiento = CATALOGO_CARGOS[c['agrupamiento']]['denominacion']
    jurisdiccion = dict(Concurso.JURISDICCIONES).get(c.get('jurisdiccion'), '')
    controles = {item['requisito']: item['fojas'] for item in datos['actuaciones_verificadas']}
    jr = {(j['representacion'], j['calidad']): j for j in datos.get('jurados', [])}
    aspirantes = datos.get('aspirantes', [])
    m = {
        '{{categoría}}': str(c['categoria']), '{{tramo}}': str(c['tramo']),
        '{{agrupamiento_nombre}}': agrupamiento, '{{agrupamiento_num}}': str(c['agrupamiento']),
        '{{CP}}': c['codigo_presupuestario'], '{{funciones}}': c['funciones'] or c['denominacion'],
        '{{unidad}}': c['unidad'], '{{dependiente_de}}': c['dependencia'],
        '{{dependencia}}': c['cadena_jerarquica'], '{{jurisdiccion}}': jurisdiccion,
        '{{el_la_unidad}}': c['articulo_unidad'],
        '{{del-de la}}': c['articulo_dependencia'] if c['articulo_dependencia'] in ('del', 'de la') else 'de',
        '{{num_expediente}}': c['expediente'],
        '{{tratamiento_iniciador_expediente}}': c['iniciador_tratamiento'],
        '{{iniciador_expediente}}': c['iniciador_nombre'],
        '{{tipo_concurso}}': dict(Concurso._meta.get_field('modalidad').choices).get(c['modalidad'], c['modalidad']),
        '{{res_convocatoria}}': c.get('resolucion_convocatoria', ''),
        '{{fojas_acta_inscripcion}}': controles.get('acta_inscriptos', ''),
        '{{fojas_acta_antecedentes}}': controles.get('valoracion_antecedentes', ''),
        '{{fojas_acta_sorteo_de_temas}}': controles.get('sorteo_temas', ''),
        '{{fojas_acta_oposicion}}': controles.get('acta_oposicion', ''),
        '{{fojas_dictamen}}': controles.get('dictamen', ''),
        '{{foja_afectación_presup}}': controles.get('afectacion_presupuestaria', ''),
        '{{res_dictamen}}': x.get('res_dictamen', ''),
        '{{fecha_vacancia_del_cargo}}': _fecha(c.get('fecha_baja')),
        '{{ex_titular}}': c.get('ex_titular', ''),
        '{{res_de_vacancia}}': c.get('resolucion_baja') or c.get('acto_vacancia', ''),
        '{{cantidad_postulantes_letras}}': _cantidad_letras(len(aspirantes)),
        '{{postulante/s}}': 'postulante' if len(aspirantes) == 1 else 'postulantes',
        '{{el/los}}': 'el' if len(aspirantes) == 1 else 'los',
        '{{presente/s}}': 'presente' if len(aspirantes) == 1 else 'presentes',
        '{{inscripto/s}}': 'inscripto' if len(aspirantes) == 1 else 'inscriptos',
        '{{fecha_sorteo_tema}}': _fecha(x.get('fecha_sorteo_tema')),
        '{{hora_sorteo_tema}}': str(x.get('hora_sorteo_tema', ''))[:5],
        '{{fecha_oposicion}}': _fecha(x.get('fecha_oposicion')),
        '{{hora_oposicion}}': str(x.get('hora_oposicion', ''))[:5],
        '{{tratamiento_autoridad_asistente}}': x.get('autoridad_tratamiento', ''),
        '{{cargo_autoridad_asistente}}': x.get('autoridad_cargo', ''),
        '{{nombre_autoridad_asistente_con_titulos}}': x.get('autoridad_nombre', ''),
        '{{ganador_a}}': '', '{{dni_ganador_a}}': '',
        '{{dni_ganador_a \\@ "0\'.\'000\'.\'000"}}': '',
    }
    for rep, nombre in [('dependencia', 'dependencia'), ('rectorado', 'rectorado'), ('personal', 'estamento')]:
        for calidad in ('titular', 'suplente'):
            j = jr.get((rep, calidad))
            sufijo = f'{calidad}_{nombre}'
            nombre_clave = f'{{{{jurado_{sufijo}}}}}'
            m[nombre_clave] = _nombre(j) if j else ''
            m[f'{{{{titulo_jurado_{sufijo}}}}}'] = j['titulo'] if j else ''
            m[f'{{{{tratamiento_jurado_{sufijo}}}}}'] = x.get(f'tratamiento_{sufijo}', '')
            if rep == 'dependencia':
                m[f'{{{{jurado_{calidad}_dependecia}}}}'] = _nombre(j) if j else ''
    return m


def _firma_asistentes(documento, jurados):
    # El original dispone las firmas de tres jurados en columnas tabuladas.
    p_nombre, p_rol = documento.paragraphs[19], documento.paragraphs[20]
    p_nombre.text = '\t'.join(_nombre(j) for j in jurados)
    p_rol.text = '\t'.join('Miembro del Jurado' for _ in jurados)


def _llenar_firmas_postulantes(documento, presentes):
    tabla = documento.tables[0]
    while len(tabla.rows) - 1 < len(presentes):
        tabla.add_row()
    while len(tabla.rows) - 1 > len(presentes):
        tabla._tbl.remove(tabla.rows[-1]._tr)
    for i, registro in enumerate(presentes, 1):
        tabla.cell(i, 0).text = _nombre(registro)
        tabla.cell(i, 1).text = _dni(registro)
        tabla.cell(i, 2).text = ''


def _acta(documento, datos, m):
    x = datos['complementarios']
    m['{{fecha_sorteo_tema}}'] = _fecha_acta(x.get('fecha_sorteo_tema'))
    presentes_j = [j for j in datos['jurados'] if str(j['id']) in x['jurados_presentes']]
    presentes_a = [a for a in datos['aspirantes'] if str(a['id']) in x['aspirantes_presentes']]
    veedores = [v for v in datos.get('veedores', []) if str(v['id']) == str(x.get('veedor_id', ''))]
    m['{{el/los}}'] = 'el' if len(presentes_a) == 1 else 'los'
    m['{{postulante/s}}'] = 'postulante' if len(presentes_a) == 1 else 'postulantes'
    m['{{presente/s}}'] = 'presente' if len(presentes_a) == 1 else 'presentes'
    if not presentes_j or not x['autoridad_nombre']:
        raise ValidationError('El sorteo requiere una autoridad fiscalizadora y al menos un miembro del jurado presente.')
    texto = documento.paragraphs[2].text
    lugar_jurados = re.search(r'los miembros integrantes del Jurado: .*?, designados en el Anexo de la Resolución', texto)
    if not lugar_jurados:
        raise ValidationError('Cambió el modelo original de acta de sorteo; revisá su versión.')
    asistentes = ', '.join(_jurado(j) + ' en representación de ' + {
        'dependencia': 'la Autoridad de la Dependencia Convocada',
        'rectorado': 'Rectorado', 'personal': 'el Personal de la Dependencia',
    }[j['representacion']] for j in presentes_j)
    _reemplazar(documento.paragraphs[2], {lugar_jurados.group(): f'los miembros integrantes del Jurado: {asistentes}, designados en el Anexo de la Resolución'})
    texto = documento.paragraphs[2].text
    lugar_aspirantes = re.search(r'se \{\{encuentra/n\}\}.*?\{\{nombre_postulante_5\}\} ', texto)
    if not lugar_aspirantes:
        raise ValidationError('Cambió el bloque de postulantes del modelo de sorteo.')
    representacion = ('se encuentran presentes los postulantes inscriptos, ' + _lista_personas(presentes_a)) if len(presentes_a) > 1 else (
        'se encuentra presente el postulante inscripto, ' + _lista_personas(presentes_a) if presentes_a else
        'no se registran postulantes presentes')
    _reemplazar(documento.paragraphs[2], {lugar_aspirantes.group(): representacion + ' '})
    if x['veedor_presencia'] == 'presente':
        if not veedores:
            raise ValidationError('Seleccioná el veedor presente del concurso.')
        veedor_txt = f'se encuentra presente el veedor gremial de APUNSJ, {_nombre(veedores[0])}.'
    elif x['veedor_presencia'] == 'ausente':
        veedor_txt = 'se deja constancia de la ausencia del veedor gremial de APUNSJ.'
    else:
        veedor_txt = ('se deja constancia que no se ha recepcionado información de la Asociación del Personal '
            'Nodocente de la Universidad (APUNSJ) designando veedor para el presente concurso, '
            'habiéndosele cursado invitación mediante Oficio Nº ' + x['oficio_veedor'] + '.')
    texto = documento.paragraphs[2].text
    comienzo = texto.find('se deja constancia que no se ha recepcionado información')
    if comienzo < 0:
        raise ValidationError('El modelo de sorteo cambió: falta la mención del veedor.')
    _reemplazar(documento.paragraphs[2], {texto[comienzo:].strip(): veedor_txt})
    _reemplazar_literal_obligatorio(
        documento.paragraphs[3], 'Lic. Marcelo Javier PINOS',
        _jurado(next(j for j in datos['jurados'] if str(j['id']) == x['jurado_autor_tema'])),
        'autor del temario sorteado',
    )
    _reemplazar_literal_obligatorio(
        documento.paragraphs[11], 'doce horas', x['hora_cierre'] + ' horas',
        'hora de cierre del acta',
    )
    _firma_asistentes(documento, presentes_j)
    _llenar_firmas_postulantes(documento, presentes_a)
    if not presentes_a:
        _quitar(documento.paragraphs[9])
        _quitar(documento.paragraphs[10])
        # No presentar como notificación un acta sin postulantes presentes.
        documento.tables[0]._element.getparent().remove(documento.tables[0]._element)
    m['{{sede_oposicion}}'] = x['sede_oposicion']
    m['{{caracter_oposicion}}'] = x['caracter_oposicion']
    # Primero resolver las demás referencias por índice, luego insertar listas Word.
    poner_temario(documento, 7, x['tema_practica'], num_vineta=2)
    poner_temario(documento, 5, x['tema_teoria'], num_vineta=1)


def _postulantes(documento, datos):
    aspirantes = datos['aspirantes']
    for p in list(documento.paragraphs):
        if '{{nombre_postulante_1}}' not in p.text:
            continue
        patron = re.compile(r'\{\{nombre_postulante_1\}\}.*?\{\{dni_postulante_5\}\}')
        match = patron.search(p.text)
        if not match:
            raise ValidationError('El modelo de resolución cambió: no se encontró la lista de postulantes.')
        _reemplazar(p, {match.group(): _lista_personas(aspirantes, con_dni=True)})


def _dictamen(documento, datos):
    resultados = datos['resultados']['orden_merito']
    lineas = [f'{i}°- {r["apellido"]}, {r["nombre"]}, DNI {r["dni"]}'
              for i, r in enumerate(resultados, 1)]
    if not lineas:
        raise ValidationError('No hay aspirantes con mérito acreditado para el proyecto de aprobación.')
    # En el original hay tres renglones de mérito; se amplían a N sin inventar ganadores.
    marcador = next((p for p in documento.paragraphs if '{{}}' in p.text), None)
    if marcador is None:
        raise ValidationError('El modelo de dictamen no contiene el bloque de orden de mérito.')
    _reemplazar(marcador, {'1°- {{}}': lineas[0]})
    parrafos = documento.paragraphs
    idx = next(i for i, p in enumerate(parrafos) if p._p is marcador._p)
    r2, r3 = parrafos[idx+1:idx+3]
    for p in (r2, r3):
        _quitar(p)
    from docx.oxml import OxmlElement
    for linea in lineas[1:]:
        nuevo = OxmlElement('w:p')
        marcador._element.addnext(nuevo)
        from docx.text.paragraph import Paragraph
        Paragraph(nuevo, marcador._parent).text = linea
        marcador = Paragraph(nuevo, marcador._parent)


def _considerando_vacancia_designacion(documento, datos):
    """Reemplaza solo el considerando de ejemplo, tomando la causa del concurso.

    Los otros considerandos y el texto dispositivo del modelo permanecen intactos.
    Si cambió el original, bloquea la emisión antes de dejar una jubilación ficticia.
    """
    inicio = 'Que según lo informado en las referidas actuaciones, el cargo se encuentra vacante'
    parrafos = [p for p in documento.paragraphs if p.text.startswith(inicio)]
    if len(parrafos) != 1 or not all(frase in parrafos[0].text for frase in (
        'beneficio jubilatorio', 'renuncia definitiva fue aceptada',
    )):
        raise ValidationError(
            'Cambió el considerando de vacancia del modelo de designación; '
            'revisá el archivo fuente antes de generar.'
        )
    texto_actual = parrafos[0].text
    texto_vacancia = causa_de_vacancia(datos)
    _reemplazar_literal_obligatorio(
        parrafos[0], texto_actual, texto_vacancia, 'causa de vacancia',
    )


def _designacion(documento, datos, m):
    _considerando_vacancia_designacion(documento, datos)
    x = datos['complementarios']
    elegido = next((a for a in datos['aspirantes'] if str(a['id']) == str(x.get('aspirante_id'))), None)
    if elegido is None:
        raise ValidationError('Seleccioná un aspirante registrado para la designación.')
    m['{{ganador_a}}'] = _nombre(elegido)
    m['{{dni_ganador_a}}'] = _dni(elegido)
    m['{{dni_ganador_a \\@ "0\'.\'000\'.\'000"}}'] = _dni(elegido)
    bloques_limitacion = [p for p in documento.paragraphs if
        p.text.startswith('Que corresponde limitar,') or p.text.startswith('ARTÍCULO 2º.- Limitar')]
    if len(bloques_limitacion) != 2:
        raise ValidationError('Cambió el modelo de designación: revisá los dos párrafos de limitación anterior.')
    if x.get('aplica_limitacion'):
        requeridos = ('resolucion_anterior', 'categoria_anterior', 'tramo_anterior',
                      'agrupamiento_anterior', 'cp_anterior', 'funciones_anteriores', 'unidad_anterior')
        if any(not x.get(k) for k in requeridos):
            raise ValidationError('Completá los datos de la designación anterior que corresponde limitar.')
        cambios_previos = {
            '3164/23-R': x['resolucion_anterior'],
            'Categoría 7': 'Categoría ' + x['categoria_anterior'],
            'Tramo Inicial': 'Tramo ' + x['tramo_anterior'],
            'Administrativo 01': x['agrupamiento_anterior'],
            '11991': x['cp_anterior'],
            'Departamento de Rendición de Cuentas': x['funciones_anteriores'],
            'Dirección de Contabilidad de la Dirección General de Contabilidad y Finanzas (SAF del Rectorado)': x['unidad_anterior'],
        }
        for ejemplo in cambios_previos:
            if not any(ejemplo in p.text for p in bloques_limitacion):
                raise ValidationError(
                    'Cambió el modelo institucional de designación: '
                    f'no se encontró el dato de ejemplo «{ejemplo}» para sustituir.'
                )
        for p in bloques_limitacion:
            _reemplazar(p, cambios_previos)
    else:
        for p in bloques_limitacion:
            _quitar(p)
        for p in documento.paragraphs:
            if p.text.startswith('ARTÍCULO 3º.-'):
                _reemplazar(p, {'ARTÍCULO 3º.-': 'ARTÍCULO 2º.-'})
            elif p.text.startswith('ARTICULO 4º.-'):
                _reemplazar(p, {'ARTICULO 4º.-': 'ARTICULO 3º.-'})
            elif p.text.startswith('ARTÍCULO 5º.-'):
                _reemplazar(p, {'ARTÍCULO 5º.-': 'ARTÍCULO 4º.-'})
    # El cuadro presupuestario original contiene datos de ejemplo; se sustituyen las 14 posiciones.
    tabla = documento.tables[0]
    campos = ('ejercicio', 'grupo', 'ppal', 'subppal', 'subsubppal', 'fte', 'fte',
              'prog', 'sprog', 'proy', 'activ', 'obra', 'finalidad', 'funcion')
    for j, clave in enumerate(campos):
        tabla.cell(6, j).text = str(x['presupuesto_' + clave])
    tabla.cell(1, 3).text = str(x['presupuesto_unidad'])
    tabla.cell(2, 3).text = str(x['presupuesto_unidad_cargo'])
    # Datos descriptivos del presupuesto: ubicaciones originales, sin modificar otros textos.
    for fila, clave in [(8, 'planta'), (9, 'inciso'), (10, 'partida'), (11, 'caracter'), (12, 'escalafon')]:
        if fila < len(tabla.rows):
            tabla.cell(fila, 3).text = str(x['presupuesto_'+clave])


def renderizar_fuente(tipo, datos):
    """El DOCX fuente es el documento resultante, sin reconstrucción de considerandos."""
    path = Path(settings.BASE_DIR) / 'documentos' / 'modelos' / FUENTES_WORD[tipo.codigo]
    if not path.is_file():
        raise ValidationError('No se encuentra el modelo Word original.')
    documento = Document(path)
    m = _comunes(datos)
    if tipo.codigo == 'sorteo_acta':
        _acta(documento, datos, m)
    else:
        _postulantes(documento, datos)
        if tipo.codigo == 'dictamen':
            _dictamen(documento, datos)
        else:
            _designacion(documento, datos, m)
    for p in documento.paragraphs:
        _reemplazar(p, m)
    for t in documento.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    _reemplazar(p, m)
    for p in documento.paragraphs:
        if PATRON.search(p.text):
            raise ValidationError('El modelo contiene variables sin resolver: ' + PATRON.search(p.text).group())
    for t in documento.tables:
        for fila in t.rows:
            for celda in fila.cells:
                if PATRON.search(celda.text):
                    raise ValidationError('El modelo contiene variables sin resolver en el presupuesto o firmas.')
    output = BytesIO()
    documento.save(output)
    return output.getvalue()
