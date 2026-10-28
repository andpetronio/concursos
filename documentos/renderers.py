"""Composición Word editable; sin reemplazos de cadenas ni dependencias de red."""
from io import BytesIO
import re
from django.core.exceptions import ValidationError
from docx import Document
from docx.shared import Cm, Pt


def nombre_completo(persona):
    return " ".join(persona[campo].strip() for campo in ("titulo", "nombre", "apellido") if persona[campo].strip())


def documento_base():
    documento = Document()
    seccion = documento.sections[0]
    seccion.top_margin = seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = seccion.right_margin = Cm(2.5)
    normal = documento.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(8)
    seccion.header.paragraphs[0].text = "Universidad Nacional de San Juan · Dirección de Concursos Nodocentes"
    seccion.footer.paragraphs[0].text = "BORRADOR NO EMITIDO · Requiere revisión administrativa"
    documento.add_paragraph("BORRADOR NO EMITIDO", "Heading 1")
    return documento


def oficio_apunsj(documento, datos, *, membrete=None):
    """Texto institucional actualizado: datos del concurso sin hechos de ejemplo."""
    d = datos["destinatario"]
    c = datos["concurso"]
    x = datos["complementarios"]
    from concursos.catalogos import CATALOGO_CARGOS
    agrupamiento = CATALOGO_CARGOS[c["agrupamiento"]]["denominacion"]
    tratamiento = "Señor" if d["tratamiento"] == "Sr." else "Señora"
    _encabezado_oficio(documento, datos, (
        tratamiento,
        f'{d["cargo"]} de la',
        "Asociación del Personal Nodocente de la UNSJ",
        nombre_completo(d),
    ), membrete)
    unidad = grupo_con_articulo(c, "unidad", "articulo_unidad", "en")
    dependencia = grupo_con_articulo(c, "dependencia", "articulo_dependencia", "de")
    _cuerpo_oficio(documento,
        "Me dirijo a usted con el propósito de solicitarle tenga a bien comunicar la "
        "designación de un veedor en representación de la Asociación que preside, "
        "conforme a lo previsto en el artículo 25º de la Ordenanza Nº 33/10-CS "
        "y en el artículo 31º del Convenio Colectivo de Trabajo para el Personal "
        "Nodocente, homologado por Decreto Nº 366/06-PEN.")
    _cuerpo_oficio(documento,
        f'Dicha intervención corresponde al Concurso {modalidad(datos)} de antecedentes '
        f'y oposición que se tramita por Expediente Nº {c["expediente"]}, destinado a '
        'cubrir, con carácter efectivo, un cargo de Planta Permanente del Personal '
        f'Nodocente, Categoría {c["categoria"]}, Tramo {c["tramo"]}, '
        f'Agrupamiento {c["agrupamiento"]}: {agrupamiento}, Código Presupuestario Nº '
        f'{c["codigo_presupuestario"]}, con funciones de “{c["denominacion"]}” '
        f'{unidad}, dependiente {dependencia} de la Universidad Nacional de San Juan, '
        'conforme a las disposiciones de la Ordenanza Nº 33/10-CS y su modificatoria '
        'Nº 20/11-CS.')
    _cuerpo_oficio(documento,
        'A tal efecto, solicito se informe el nombre y apellido, DNI, número de teléfono '
        'y correo electrónico de la persona designada, así como una dirección de correo '
        'electrónico institucional a la cual remitir copia de la Resolución '
        f'Nº {x["res_convocatoria"]}.')
    _cuerpo_oficio(documento,
        'En virtud de la necesidad de dar continuidad a las actuaciones administrativas '
        'del concurso, agradeceré que la información solicitada sea remitida a la mayor '
        'brevedad posible.')
    _cuerpo_oficio(documento, 'Sin otro particular, saludo a usted atentamente.')
    # No presumir firmante ni insertar firma o sellos; queda espacio para suscripción ológrafa.
    espacio = documento.add_paragraph()
    espacio.paragraph_format.space_after = Pt(28)



RENDERERS = {"apunsj": oficio_apunsj}


INVALIDOS_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def validar_textos_xml(valor):
    if isinstance(valor, str) and INVALIDOS_XML.search(valor):
        raise ValidationError("Los datos contienen caracteres de control no admitidos en Word. "
                              "Revisá los textos del concurso y del destinatario antes de generar.")
    if isinstance(valor, dict):
        for clave, dato in valor.items():
            validar_textos_xml(clave)
            validar_textos_xml(dato)
    elif isinstance(valor, (list, tuple)):
        for dato in valor:
            validar_textos_xml(dato)


def renderizar(tipo, instantanea, *, membrete=None):
    validar_textos_xml(instantanea)
    renderer = RENDERERS.get(tipo.codigo)
    if renderer is None:
        raise ValidationError("Plantilla documental no disponible.")
    documento = base_institucional(tipo.codigo)
    if tipo.codigo in ("jurados", "elecciones", "apunsj"):
        renderer(documento, instantanea, membrete=membrete)
    else:
        renderer(documento, instantanea)
    if tipo.codigo in ("elecciones", "apunsj"):
        _uniformar_tipografia_oficios(documento)
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()

# Reconstrucción de los cinco originales. Se preserva su estructura, no sus hechos de ejemplo.
# Preset: standard_business_brief; override institucional PDF: Arial 12, A4,
# márgenes 25/15/25/30 mm, interlineado 1,5, títulos negros centrados.
# Cartas: membrete institucional sin firma ni conmemoración anual; notificación Letter.
from pathlib import Path
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Twips, Emu
from docx.enum.style import WD_STYLE_TYPE
from participantes.models import Jurado
from . import convocatoria_textos as ctext


def base_institucional(tipo):
    documento = Document()
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = (Cm(21.59), Cm(27.94)) if tipo in ("jurados", "elecciones", "apunsj") else (Cm(21), Cm(29.7))
    seccion.top_margin = Cm(5) if tipo in ("jurados", "elecciones", "apunsj") else Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin, seccion.right_margin = Cm(3), Cm(1.5)
    seccion.header_distance = seccion.footer_distance = Cm(1.25)
    for nombre in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Header", "Footer"):
        estilo = documento.styles[nombre]
        estilo.font.name, estilo.font.size = "Arial", Pt(12)
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        formato = estilo.paragraph_format
        formato.space_before, formato.space_after = Pt(0), Pt(6)
        formato.line_spacing = 1.5
    if tipo in ("jurados", "elecciones", "apunsj"):
        documento.styles["Normal"].font.size = Pt(11)
        documento.styles["Normal"].paragraph_format.line_spacing = 1.15
    normal = documento.styles["Normal"].paragraph_format
    normal.alignment, normal.first_line_indent = WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(2)
    for nombre in ("Encabezado institucional", "Metadato institucional", "Celda presupuestaria", "Sección anexo"):
        estilo = documento.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        estilo.base_style = documento.styles["Normal"]
        estilo.paragraph_format.first_line_indent = Cm(0)
        estilo.paragraph_format.space_before = Pt(0)
        estilo.paragraph_format.space_after = Pt(3)
        estilo.paragraph_format.line_spacing = 1
    documento.styles["Encabezado institucional"].font.bold = True
    documento.styles["Encabezado institucional"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    documento.styles["Sección anexo"].font.bold = True
    documento.styles["Sección anexo"].paragraph_format.keep_with_next = True
    # El anexo contiene hasta doce apartados extensos: tipografía compacta solo allí,
    # sin alterar el cuerpo de los artículos resolutivos.
    anexo_detalle = documento.styles.add_style("Detalle anexo", WD_STYLE_TYPE.PARAGRAPH)
    anexo_detalle.base_style = documento.styles["Normal"]
    anexo_detalle.font.size = Pt(10)
    anexo_detalle.paragraph_format.first_line_indent = Cm(0)
    anexo_detalle.paragraph_format.space_after = Pt(2)
    anexo_detalle.paragraph_format.line_spacing = 1.08
    documento.styles["Celda presupuestaria"].font.size = Pt(7)
    documento.styles["Celda presupuestaria"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # El proyecto de convocatoria se identifica como tal en su encabezado;
    # no agrega advertencias administrativas al pie de página.
    if tipo not in ("convocatoria", "jurados", "elecciones", "apunsj"):
        pie = seccion.footer.paragraphs[0]
        pie.text = "BORRADOR NO EMITIDO · Requiere revisión administrativa · Sin firma"
        pie.paragraph_format.first_line_indent = Cm(0)
        pie.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pie.paragraph_format.line_spacing = 1
        for run in pie.runs:
            run.font.size = Pt(8)
    return documento


def _uniformar_tipografia_oficios(documento):
    """Evita sustituciones de fuente entre fecha, destinatario, cuerpo y firma.

    El membrete es una imagen aportada por la institución; su tipografía interna
    no puede alterarse sin reemplazar la propia imagen.
    """
    familia = "Arial"
    for nombre in ("Normal", "Metadato institucional"):
        estilo = documento.styles[nombre]
        estilo.font.name = familia
        estilo.font.size = Pt(11)
        fuentes = estilo._element.get_or_add_rPr().rFonts
        if fuentes is not None:
            fuentes.set(qn("w:eastAsia"), familia)
            fuentes.set(qn("w:cs"), familia)
    for parrafo_oficio in documento.paragraphs:
        for fragmento in parrafo_oficio.runs:
            if not fragmento.text:
                continue
            fragmento.font.name = familia
            fragmento.font.size = Pt(11)
            fuentes = fragmento._element.get_or_add_rPr().rFonts
            if fuentes is not None:
                fuentes.set(qn("w:eastAsia"), familia)
                fuentes.set(qn("w:cs"), familia)


def parrafo(documento, texto, *, etiqueta=None, estilo=None):
    p = documento.add_paragraph(style=estilo)
    if etiqueta:
        p.add_run(etiqueta + " ").bold = True
    p.add_run(texto)
    return p


def valor_presentado(datos, campo):
    return dict(datos["concurso"]["datos_presentados"]).get(campo, "")


def cargo(datos):
    c = datos["concurso"]
    return (f'{c["denominacion"]}, categoría {c["categoria"]}, tramo {c["tramo"]}, '
            f'agrupamiento {valor_presentado(datos, "Agrupamiento")} ({c["agrupamiento"]}), '
            f'código presupuestario {c["codigo_presupuestario"]}, unidad {c["unidad"]}, '
            f'dependencia {c["dependencia"]}, {c["cadena_jerarquica"]}')


def modalidad(datos):
    from concursos.models import Concurso
    return dict(Concurso._meta.get_field("modalidad").choices)[datos["concurso"]["modalidad"]]


def persona_texto(persona):
    dni = persona["dni"]
    if re.fullmatch(r"[0-9]{8}", dni):
        dni = f"{dni[:2]}.{dni[2:5]}.{dni[5:]}"
    return f'{persona["nombre"]} {persona["apellido"]}, DNI {dni}'


def jurados_texto(datos):
    representaciones = dict(Jurado.Representacion.choices)
    return "; ".join(f'{j["calidad"]} en representación de {representaciones[j["representacion"]]}: '
                     f'{j["titulo"]} {persona_texto(j["persona"])}, unidad {j["unidad"]}'
                     for j in datos["jurados"])


def foja(datos, codigo):
    return next(a["fojas"] for a in datos["actuaciones_verificadas"] if a["requisito"] == codigo)


def antecedente_iniciador(datos):
    c = datos["concurso"]
    nombre = " ".join(v.strip() for v in (
        c.get("iniciador_tratamiento") or "", c.get("iniciador_nombre") or "") if v.strip())
    cargo = (c.get("iniciador_cargo") or "").strip()
    if not nombre and not cargo:
        return ""
    return f"{nombre}, {cargo}" if nombre and cargo else nombre or cargo


def antecedente_vacancia(datos):
    c = datos["concurso"]
    fecha = c.get("fecha_baja")
    if fecha and len(fecha) == 10 and fecha[4] == "-":
        anio, mes, dia = fecha.split("-")
        fecha = f"{dia}/{mes}/{anio}"
    etiquetas = (("ex titular", c.get("ex_titular")),
                 ("motivo de baja", c.get("motivo_baja")),
                 ("fecha de baja", fecha),
                 ("resolución de baja", c.get("resolucion_baja")))
    partes = [f"{nombre}: {valor}" for nombre, valor in etiquetas if valor]
    return "Antecedentes registrados de vacancia: " + "; ".join(partes) + "." if partes else ""


def apertura_resolucion(documento, datos, titulo, *, marco_normativo=None):
    x, c = datos["complementarios"], datos["concurso"]
    parrafo(documento, "PROYECTO DE RESOLUCIÓN", estilo="Encabezado institucional")
    parrafo(documento, titulo, estilo="Encabezado institucional")
    for etiqueta, valor in (("Categoría", c["categoria"]), ("Tramo", c["tramo"]),
        ("Agrupamiento", f'{valor_presentado(datos, "Agrupamiento")} ({c["agrupamiento"]})'),
        ("Código Presupuestario", c["codigo_presupuestario"]),
        ("Funciones", c["funciones"]),
        ("Unidad", c["unidad"]), ("Dependencia", f'{c["dependencia"]} / {c["cadena_jerarquica"]}')):
        if etiqueta == "Funciones" and not valor.strip():
            continue
        parrafo(documento, valor, etiqueta=etiqueta + ":", estilo="Metadato institucional")
    if x.get("lugar_fecha"):
        parrafo(documento, x["lugar_fecha"], estilo="Metadato institucional")
    parrafo(documento, "VISTO:", estilo="Sección anexo")
    iniciador = antecedente_iniciador(datos)
    detalle_inicio = f", iniciado por {iniciador}" if iniciador else ""
    parrafo(documento, f'El Expediente N° {c["expediente"]}{detalle_inicio}, '
             f'relativo al concurso {modalidad(datos)} para cubrir el cargo de {cargo(datos)}; y')
    parrafo(documento, "CONSIDERANDO:", estilo="Sección anexo")
    parrafo(documento, x.get("marco_normativo") or marco_normativo or ctext.MARCO_NORMATIVO)


def resuelve(documento, datos):
    parrafo(documento, "Por ello y en uso de sus atribuciones", estilo="Encabezado institucional")
    parrafo(documento, datos["complementarios"].get("autoridad_resuelve") or "EL RECTOR DE LA UNIVERSIDAD NACIONAL DE SAN JUAN", estilo="Encabezado institucional")
    parrafo(documento, "RESUELVE:", estilo="Encabezado institucional")


def articulo(documento, numero, texto, *, sin_sangria=False):
    # Etiquetas jurídicas, no una lista simulada: los números integran cada disposición.
    p = parrafo(documento, texto, etiqueta=f"ARTÍCULO {numero}°.-")
    if sin_sangria:
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
    return p


def cierre_resolucion(documento, datos, numero):
    articulo(documento, numero, "Regístrese, comuníquese y oportunamente archívese.")
    parrafo(documento, "Resolución N°: no asignado; proyecto sin emisión.", estilo="Metadato institucional")
    if datos["complementarios"].get("firmante"):
        parrafo(documento, "Autoridad firmante propuesta (sin firma): " + datos["complementarios"]["firmante"],
                estilo="Metadato institucional")


def tabla_fija(documento, filas, anchos):
    tabla = documento.add_table(rows=len(filas), cols=len(anchos))
    tabla.style = "Table Grid"
    tabla.autofit = False
    props = tabla._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), str(sum(anchos)))
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "60")
    indent.set(qn("w:type"), "dxa")
    props.append(indent)
    margenes = OxmlElement("w:tblCellMar")
    for lado in ("top", "left", "bottom", "right"):
        margen = OxmlElement(f"w:{lado}")
        margen.set(qn("w:w"), "60")
        margen.set(qn("w:type"), "dxa")
        margenes.append(margen)
    props.append(margenes)
    for columna, ancho in zip(tabla.columns, anchos):
        columna.width = Twips(ancho)
    for fila, textos in zip(tabla.rows, filas):
        for celda, texto, ancho in zip(fila.cells, textos, anchos):
            celda.width = Twips(ancho)
            celda.text = str(texto)
            for p in celda.paragraphs:
                p.style = documento.styles["Celda presupuestaria"]
    return tabla


def presupuesto(documento, datos, *, para_convocatoria=False):
    x, c = datos["complementarios"], datos["concurso"]
    if "presupuesto_unidad" in x:  # La designación conserva su formulario propio.
        valores = {nombre: x[f"presupuesto_{nombre}"] for nombre in (
            "unidad", "unidad_cargo", "ejercicio", "grupo", "ppal", "subppal",
            "subsubppal", "fte", "prog", "sprog", "proy", "activ", "obra",
            "finalidad", "funcion", "planta", "inciso", "partida", "caracter", "escalafon", "numero_cargo")}
    else:
        control = next(a for a in datos["actuaciones_verificadas"]
            if a["requisito"] == "afectacion_presupuestaria")
        valores = {**ctext.PRESUPUESTO_FIJO,
            "unidad": jurisdiccion_texto(datos) if para_convocatoria else ctext.PRESUPUESTO_FIJO["unidad"],
            "unidad_cargo": f'{c["unidad"]} - {c["dependencia"]} - {c["cadena_jerarquica"]}',
            "ppal": control["partida_principal"], "activ": control["actividad_presupuestaria"],
            "numero_cargo": c["codigo_presupuestario"]}
    def valor(nombre):
        return valores[nombre]
    ancho = Emu(documento.sections[0].page_width - documento.sections[0].left_margin - documento.sections[0].right_margin).twips
    parrafo(documento, "PARTE DE AFECTACIÓN - PERSONAL NODOCENTE", estilo="Encabezado institucional")
    tabla_fija(documento, [("UNIDAD", valor("unidad")), ("UNIDAD DONDE SE UBICA EL CARGO", valor("unidad_cargo"))],
               [ancho // 3, ancho - ancho // 3])
    claves = ("ejercicio", "grupo", "ppal", "subppal", "subsubppal", "fte", "prog", "sprog", "proy", "activ", "obra", "finalidad", "funcion")
    titulos = ("EJERCICIO\nFINANCIERO", "GRUPO\nPRESU.", "UNIDAD", "", "", "FTE", "RED PROGRAMÁTICA", "", "", "", "", "FINALIDAD Y FUNCIÓN", "")
    subtitulos = ("", "", "PPAL", "SUB\nPPAL", "SUB SUB\nPPAL", "", "PROG", "SPROG", "PROY", "ACTIV", "OBRA", "FINALIDAD", "FUNC.")
    anchos = [ancho // 13] * 13
    anchos[-1] += ancho - sum(anchos)
    tabla = tabla_fija(documento, [titulos, subtitulos, [valor(k) for k in claves]], anchos)
    for inicio, fin in ((2, 4), (6, 10), (11, 12)):
        tabla.cell(0, inicio).merge(tabla.cell(0, fin))
    for fila in tabla.rows[:2]:
        repetir = OxmlElement("w:tblHeader")
        fila._tr.get_or_add_trPr().append(repetir)
    detalles = [("PLANTA", valor("planta")), ("INCISO", valor("inciso")), ("PARTIDA PPAL", valor("partida")),
        ("CARÁCTER DESIGNACIÓN", valor("caracter")), ("ESCALAFÓN", valor("escalafon")),
        ("CATEGORÍA / AGRUPAMIENTO", f'{c["categoria"]} / {c["agrupamiento"]}'),
        ("NRO. CARGO", valor("numero_cargo")), ("CÓDIGO PRESUPUESTARIO", c["codigo_presupuestario"])]
    if para_convocatoria:
        detalles = [(etiqueta, f'{c["categoria"]} / {agrupamiento_formateado(datos)}')
                    if etiqueta == "CATEGORÍA / AGRUPAMIENTO" else (etiqueta, valor)
                    for etiqueta, valor in detalles]
        detalles.append(("FUNCIONES", c["denominacion"]))
    elif c["funciones"].strip():
        detalles.append(("FUNCIONES", c["funciones"]))
    tabla_fija(documento, detalles, [ancho // 3, ancho - ancho // 3])


def numeracion_anexo(documento):
    """Numeración Word real: reinicia en 1 y conserva el contador de doce apartados."""
    raiz = documento.part.numbering_part.element
    abstract_id = max([int(n.get(qn("w:abstractNumId"))) for n in raiz.findall(qn("w:abstractNum"))] + [0]) + 1
    numero_id = max([int(n.get(qn("w:numId"))) for n in raiz.findall(qn("w:num"))] + [0]) + 1
    abstracto = OxmlElement("w:abstractNum")
    abstracto.set(qn("w:abstractNumId"), str(abstract_id))
    nivel = OxmlElement("w:lvl")
    nivel.set(qn("w:ilvl"), "0")
    for tag, valor in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1.-"), ("lvlJc", "left")):
        elem = OxmlElement(f"w:{tag}")
        elem.set(qn("w:val"), valor)
        nivel.append(elem)
    formato = OxmlElement("w:pPr")
    sangria = OxmlElement("w:ind")
    sangria.set(qn("w:left"), "420")
    sangria.set(qn("w:hanging"), "420")
    formato.append(sangria)
    nivel.append(formato)
    abstracto.append(nivel)
    raiz.append(abstracto)
    numero = OxmlElement("w:num")
    numero.set(qn("w:numId"), str(numero_id))
    referencia = OxmlElement("w:abstractNumId")
    referencia.set(qn("w:val"), str(abstract_id))
    numero.append(referencia)
    raiz.append(numero)
    return numero_id


def agrupamiento_formateado(datos):
    from concursos.catalogos import CATALOGO_CARGOS
    codigo = datos["concurso"]["agrupamiento"]
    denominacion = CATALOGO_CARGOS[codigo]["denominacion"]
    return f"{denominacion} - {codigo}"


def jurisdiccion_texto(datos):
    from concursos.models import Concurso
    codigo = datos["concurso"]["jurisdiccion"]
    return dict(Concurso.JURISDICCIONES)[codigo]


def articulo_normalizado(articulo, preposicion=""):
    """Las opciones visibles son el/la/del/de la/sin artículo; el documento contrae."""
    base = {"del": "el", "de la": "la"}.get(articulo or "", articulo or "")
    if preposicion == "de":
        return "del" if base == "el" else "de la" if base == "la" else "de"
    if preposicion == "en":
        return "en " + base if base else "en"
    return base


def grupo_con_articulo(c, campo, articulo, preposicion=""):
    return " ".join(v for v in (articulo_normalizado(c.get(articulo), preposicion), c[campo]) if v).strip()


def estructura(datos):
    c = datos["concurso"]
    return (f'{grupo_con_articulo(c, "unidad", "articulo_unidad", "en")}, '
            f'dependiente {grupo_con_articulo(c, "dependencia", "articulo_dependencia", "de")} - '
            f'{grupo_con_articulo(c, "cadena_jerarquica", "articulo_superior")} '
            f'del {jurisdiccion_texto(datos)}')


def fecha_es(valor):
    from datetime import date
    if not valor:
        return ""
    fecha = date.fromisoformat(valor)
    meses = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
             "septiembre", "octubre", "noviembre", "diciembre")
    return f"{fecha.day} de {meses[fecha.month-1]} de {fecha.year}"


def fecha_inscripcion(datos):
    return fecha_es(datos["complementarios"]["fecha_apertura"])


def causa_de_vacancia(datos):
    c = datos["concurso"]
    causa = c["causa_vacancia"]
    fecha = f' a partir del {fecha_es(c["fecha_baja"])}' if c.get("fecha_baja") else ""
    ex = c.get("ex_titular") or ""
    acto = c.get("acto_vacancia") or c.get("resolucion_baja") or ""
    res = f" mediante Resolución N.º {c['resolucion_baja']}" if c.get("resolucion_baja") else (
        f" conforme al acto {c['acto_vacancia']}" if c.get("acto_vacancia") else "")
    if causa == "nuevo":
        return ("Que el cargo objeto del presente concurso es de nueva creación" +
                (f", conforme al acto {acto}" if acto else "") + ", encontrándose disponible para su cobertura.")
    if causa == "jubilacion":
        return f"Que el cargo se encuentra vacante{fecha}, en virtud de que su ex titular, {ex}, obtuvo el beneficio jubilatorio, constando su baja{res}."
    if causa == "renuncia":
        return f"Que el cargo se encuentra vacante{fecha}, en virtud de que su ex titular, {ex}, presentó su renuncia, aceptada{res}."
    if causa == "fallecimiento":
        return f"Que el cargo se encuentra vacante{fecha}, como consecuencia del fallecimiento de quien fuera su titular, {ex}."
    if causa == "concurso":
        return f"Que el cargo se encuentra vacante{fecha}, en virtud de la designación efectiva de su anterior titular, {ex}, en otro cargo por concurso{res}."
    if causa == "efectivizacion":
        return f"Que el cargo se encuentra vacante{fecha}, en virtud de la efectivización de su anterior titular, {ex}, en otro cargo{res}."
    if causa == "otra":
        return f"Que el cargo se encuentra vacante{fecha}. {c['detalle_vacancia'].strip()}"
    raise ValidationError("Seleccioná una causa de vacancia antes de generar convocatoria.")


def jurado_por(datos, representacion, calidad):
    return next(j for j in datos["jurados"] if j["representacion"] == representacion and j["calidad"] == calidad)


def jurado_presentado(j):
    persona = j["persona"]
    titulo = (j.get("titulo") or "").strip()
    dni = persona["dni"].replace(".", "")
    if dni.isdigit() and len(dni) == 8:
        dni = f"{dni[:2]}.{dni[2:5]}.{dni[5:]}"
    return f'{titulo + " " if titulo else ""}{persona["nombre"]} {persona["apellido"]} D.N.I. N.º {dni}'


def parrafo_anexo(documento, texto):
    return parrafo(documento, texto, estilo="Detalle anexo")


def item_anexo(documento, num_id, titulo, lineas):
    p = parrafo(documento, titulo, estilo="Sección anexo")
    # Separar visualmente los apartados sin insertar párrafos en blanco.
    p.paragraph_format.space_before = Pt(9)
    props = p._p.get_or_add_pPr()
    numeros = OxmlElement("w:numPr")
    for nombre, valor in (("ilvl", "0"), ("numId", str(num_id))):
        elem = OxmlElement(f"w:{nombre}")
        elem.set(qn("w:val"), valor)
        numeros.append(elem)
    props.append(numeros)
    for linea in lineas:
        parrafo_anexo(documento, linea)


def anexo_convocatoria(documento, datos):
    c = datos["concurso"]
    documento.add_page_break()
    parrafo(documento, "ANEXO", estilo="Encabezado institucional")
    num_id = numeracion_anexo(documento)
    turno = {"manana": "7 a 14", "tarde": "14 a 21"}[c["turno_trabajo"]]
    unidad = c["unidad"]
    juris = jurisdiccion_texto(datos)
    item_anexo(documento, num_id, "CARGO A CONCURSAR:", [
        f'Categoría: {c["categoria"]}', f'Tramo: {c["tramo"]}',
        f'Agrupamiento: {agrupamiento_formateado(datos)}',
        f'Código Presupuestario: {c["codigo_presupuestario"]}',
        "Remuneración: La establecida por la Paritaria Nacional en un todo de acuerdo al Convenio Colectivo de Trabajo del Personal de Apoyo Universitario, homologado por el Decreto 366/06-PEN para la referida categoría.",
    ])
    item_anexo(documento, num_id, "UNIDAD A LA QUE PERTENECE:", [unidad])
    item_anexo(documento, num_id, "DEPENDENCIA:", [f'{c["dependencia"]} del {juris}.'])
    item_anexo(documento, num_id, "TIPO DE CONCURSO:", [
        f'{modalidad(datos)} (Ord. 33/10-CS)', "Mecanismo del Concurso: Concurso de antecedentes y prueba de oposición."])
    item_anexo(documento, num_id, "FUNCIONES:", [f'“{c["denominacion"]}”'])
    item_anexo(documento, num_id, "LUGAR DE TRABAJO:", [
        f'{unidad}. Domicilio: {c["domicilio_trabajo"]} - San Juan.', f'Horario: {turno} hs.'])
    if c["modalidad"] == "cerrado_interno":
        generales = ["7.1. REQUISITOS Y CONDICIONES GENERALES:",
            "1- Que registre una antigüedad mínima de un (1) año en la planta permanente del Personal de Apoyo Universitario.",
            "2- Que el cargo a cubrir sea de igual o superior Categoría al que revista en planta permanente el aspirante."]
    else:
        generales = ["7.1. REQUISITOS Y CONDICIONES GENERALES:",
            "[REVISAR Y COMPLETAR LOS REQUISITOS GENERALES DE ESTA MODALIDAD ANTES DE EMITIR]"]
    item_anexo(documento, num_id, "REQUISITOS, CONDICIONES GENERALES Y PARTICULARES QUE ESPECIALMENTE SE LE EXIGEN A LOS POSTULANTES:", [
        *generales, "7.2. REQUISITOS Y CONDICIONES PARTICULARES:", "1- REQUISITO MÍNIMO:", "1. Título Secundario"])
    lineas_jurados = []
    for indice, (rep, titulo) in enumerate((
            ("rectorado", "En Representación del Rectorado"),
            ("dependencia", "En Representación de la Autoridad de la Dependencia que convoca"),
            ("personal", "En Representación del Personal de la Dependencia")), 1):
        lineas_jurados.append(f"8.{indice}. {titulo}:")
        for numero, calidad in ((1, "titular"), (2, "suplente")):
            lineas_jurados.append(f'{numero}- {calidad.capitalize()}: {jurado_presentado(jurado_por(datos, rep, calidad))}')
    item_anexo(documento, num_id, "NÓMINA DE INTEGRANTES DEL JURADO:", lineas_jurados)
    item_anexo(documento, num_id, "LUGAR, FECHA DE APERTURA Y CIERRE DE INSCRIPCIONES Y LUGAR DE PRESENTACIÓN DE ANTECEDENTES:", [
        f'9.1. Apertura de Inscripciones: a partir del {fecha_inscripcion(datos)} y por el término de CINCO (5) días hábiles en horario de 8 a 12 horas.',
        "9.2. Presentación de antecedentes: solicitud de inscripción y probanzas certificadas en: Mesa de Entradas, Salidas y Archivo – Rectorado."])
    item_anexo(documento, num_id, "LUGAR DONDE PODRÁN REQUERIRSE MAYORES DETALLES:", [
        f'{unidad}. Domicilio: {c["domicilio_trabajo"]} y en Dirección de Concursos Nodocentes - Planta Baja Edificio Central - Rectorado. Mitre 396 (Este) – Capital.'])
    item_anexo(documento, num_id, "LUGAR, FECHA Y HORA EN QUE SE REALIZARÁ LA PRUEBA DE OPOSICIÓN:", [
        f'11.1. LUGAR: {unidad} - {juris}',
        "11.2. FECHA: El Jurado fijará la fecha de realización de la prueba de oposición (Art. 37º de la Ordenanza N.º 33/10-CS). El tema sorteado deberá ser fehacientemente comunicado a los aspirantes con 72 horas de anticipación a la fecha fijada para la realización de la prueba de oposición (Art. 36º de la Ordenanza N.º 33/10-CS)."])
    item_anexo(documento, num_id, "TEMARIO GENERAL:", [" "])


def articulo_convocatoria(documento, numero, texto):
    return articulo(documento, numero, texto, sin_sangria=True)


def convocatoria(documento, datos):
    c = datos["concurso"]
    # Encabezado con los nombres sin artículos y el agrupamiento sin repetir el código.
    parrafo(documento, "PROYECTO DE RESOLUCIÓN", estilo="Encabezado institucional")
    for etiqueta, valor in (("Categoría", c["categoria"]), ("Tramo", c["tramo"]),
        ("Agrupamiento", agrupamiento_formateado(datos)), ("Código Presupuestario", c["codigo_presupuestario"]),
        ("Funciones", c["denominacion"]), ("Unidad", c["unidad"]),
        ("Dependencia", f'{c["dependencia"]} - {c["cadena_jerarquica"]}')):
        parrafo(documento, valor, etiqueta=etiqueta + ":", estilo="Metadato institucional")
    parrafo(documento, "VISTO:", estilo="Sección anexo")
    iniciador = " ".join(v for v in (
        articulo_normalizado(c.get("articulo_iniciador")), c.get("iniciador_tratamiento"), c["iniciador_nombre"])
        if v).strip()
    cargo_iniciador = f', {c["iniciador_cargo"]}' if c.get("iniciador_cargo") else ""
    parrafo(documento, f'El Expediente N° {c["expediente"]} por el cual {iniciador}{cargo_iniciador}, '
        f'solicita se convoque a concurso {modalidad(datos)}, en el marco de lo establecido por el Reglamento '
        f'de Concursos del Personal Nodocente, para cubrir un cargo Categoría {c["categoria"]} '
        f'Tramo {c["tramo"]} del Agrupamiento {agrupamiento_formateado(datos)}, '
        f'Código Presupuestario {c["codigo_presupuestario"]}, con funciones de “{c["denominacion"]}” '
        f'{estructura(datos)}, y')
    parrafo(documento, "CONSIDERANDO:", estilo="Sección anexo")
    parrafo(documento, ctext.CONSIDERANDO_1)
    parrafo(documento, ctext.CONSIDERANDO_2)
    parrafo(documento, ctext.CONSIDERANDO_3)
    parrafo(documento, f'Que a fojas {foja(datos, "jurado_dependencia")} y {foja(datos, "jurado_rectorado")} '
        'de actuaciones de referencia, se proponen los miembros que integrarán el Jurado, en calidad de titulares '
        'y suplentes, en representación de la Autoridad de la Dependencia que convoca y del Rectorado, respectivamente.')
    parrafo(documento, 'Que los miembros que integrarán el Jurado en representación del Personal de la Dependencia, '
        'de acuerdo con lo dispuesto en el Acta Paritaria N.º 2/13, han sido elegidos mediante elecciones directas '
        f'por sus pares, según documentación incorporada a fojas {foja(datos, "jurado_personal")} del referido expediente.')
    parrafo(documento, f'Que a fojas {foja(datos, "afectacion_presupuestaria")} del mismo expediente, '
        'se adjunta la afectación presupuestaria del cargo objeto del concurso.')
    parrafo(documento, causa_de_vacancia(datos))
    parrafo(documento, ctext.INTERVENCION_SAF)
    parrafo(documento, "Por ello y en uso de sus atribuciones", estilo="Encabezado institucional")
    parrafo(documento, "EL RECTOR DE LA UNIVERSIDAD NACIONAL DE SAN JUAN", estilo="Encabezado institucional")
    parrafo(documento, "RESUELVE:", estilo="Encabezado institucional")
    articulo_convocatoria(documento, 1, f'Convocar a “{modalidad(datos)}” de antecedentes y oposición para cubrir un cargo, '
        'con carácter efectivo, de planta permanente del Personal Nodocente Universitario, '
        f'Categoría {c["categoria"]} del Agrupamiento {agrupamiento_formateado(datos)}, '
        f'Tramo {c["tramo"]}, Código Presupuestario {c["codigo_presupuestario"]}, '
        f'para desempeñar funciones de “{c["denominacion"]}” {estructura(datos)} '
        'de la Universidad Nacional de San Juan, en los términos y condiciones previstos en las Ordenanzas '
        '33/10-CS y 20/11-CS, cuyos requisitos, condiciones, fecha de inscripción y otra información relevante '
        'se encuentran detallados en el ANEXO de la presente resolución.')
    articulo_convocatoria(documento, 2, f'Fijar la apertura de las inscripciones a partir del {fecha_inscripcion(datos)}, '
        'en horario de 8 a 12 horas y por el término de 5 (CINCO) días hábiles, en el cual las y los interesadas/os '
        'deberán presentar solicitud de inscripción con carácter de declaración jurada y adjuntar probanzas '
        'debidamente certificadas por autoridad competente, en Mesa de Entradas, Salidas y Archivo de Rectorado '
        '– Edificio Central de Rectorado: Planta Baja, Mitre 396 (Este), Capital. La falta de presentación '
        'de dicha documentación invalidará la inscripción al concurso.')
    articulo_convocatoria(documento, 3, 'Establecer que las copias de los antecedentes y certificaciones aportadas por las y los '
        'postulantes, autenticadas por la Mesa de Entradas, Salidas y Archivo del Rectorado, pasarán a formar '
        'parte del expediente de concurso y, por consiguiente, no se reintegrarán a sus titulares.')
    articulo_convocatoria(documento, 4, ctext.PUBLICIDAD)
    articulo_convocatoria(documento, 5, ctext.OPOSICION)
    articulo_convocatoria(documento, 6, 'Convocar mediante notificación fehaciente, que deberá constar en las actuaciones '
        'del presente concurso, a la Asociación del Personal Nodocente de la Universidad Nacional de San Juan '
        '–APUNSJ– a participar con carácter de veedor, mediante la designación de UN (1) representante.')
    articulo_convocatoria(documento, 7, f'El cargo objeto del concurso se encuentra presupuestado a fojas '
        f'{foja(datos, "afectacion_presupuestaria")} del mencionado expediente, según el siguiente detalle:')
    presupuesto(documento, datos, para_convocatoria=True)
    articulo_convocatoria(documento, 8, 'Regístrese, comuníquese y oportunamente archívese.')
    anexo_convocatoria(documento, datos)

def antecedentes_posteriores(documento, datos):
    x = datos["complementarios"]
    parrafo(documento, f'Que la convocatoria consta en la resolución oficial {x["res_convocatoria"]}.')
    parrafo(documento, f'Que obra el acta de cierre de inscripción a fs. {foja(datos, "acta_inscriptos")}. {x["antecedentes_inscriptos"]}')
    parrafo(documento, "Personas inscriptas: " + "; ".join(persona_texto(a["persona"]) for a in datos["aspirantes"]) + ".")
    parrafo(documento, "Integración del jurado: " + jurados_texto(datos) + ".")
    for campo, codigo, etiqueta in (("valoracion", "valoracion_antecedentes", "Valoración de antecedentes"),
        ("sorteo", "sorteo_temas", "Sorteo de temas"), ("prueba", "acta_oposicion", "Prueba de oposición"),
        ("resultado", "dictamen", "Dictamen")):
        parrafo(documento, f'{etiqueta}, fs. {foja(datos, codigo)}. {x[campo]}')
    for campo in ("impugnaciones", "comunicacion", "tramites"):
        parrafo(documento, x[campo])


def merito(documento, datos):
    """Dictamen ordinario: tabla de puntajes calculados; designación sigue provisional."""
    ancho = Emu(documento.sections[0].page_width - documento.sections[0].left_margin - documento.sections[0].right_margin).twips
    if datos.get("resultados"):
        resultados = datos["resultados"]
        escala = resultados["escala"]
        def fila_puntajes(registro, orden=""):
            p = registro
            dni = p["dni"]
            if dni.isdigit() and len(dni) == 8:
                dni = f"{dni[:2]}.{dni[2:5]}.{dni[5:]}"
            return (str(orden), f'{p["apellido"]}, {p["nombre"]}', dni,
                    p["antecedentes"].replace(".", ","),
                    p["oposicion"].replace(".", ","), p["total"].replace(".", ","))
        cabecera = ("ORDEN", "APELLIDO Y NOMBRE", "DNI", "ANT. /34", "OPOS. /51", "TOTAL /85")
        weights = (9, 28, 20, 14, 14, 15)
        widths = [ancho * w // sum(weights) for w in weights]
        widths[-1] = ancho - sum(widths[:-1])
        parrafo(documento, "ORDEN DE MÉRITO (mínimo 42,50 puntos)", estilo="Sección anexo")
        tabla_fija(documento, [cabecera] + [fila_puntajes(fila, fila["puesto"])
            for fila in resultados["orden_merito"]], widths)
        if resultados["puntaje_insuficiente"]:
            parrafo(documento, "ASPIRANTES QUE NO ALCANZARON EL PUNTAJE MÍNIMO", estilo="Sección anexo")
            tabla_fija(documento, [cabecera] + [fila_puntajes(fila, "—")
                for fila in resultados["puntaje_insuficiente"]], widths)
        parrafo(documento, "Escala de evaluación aplicada: antecedentes hasta 34 puntos, oposición hasta 51 puntos; máximo 85; mínimo 42,50. El rubro desempeño laboral no se considera (artículo 52 del Anexo de la Ordenanza 33/10-CS).")
        return
    aspirantes = {a["id"]: a["persona"] for a in datos["aspirantes"]}
    parrafo(documento, "ORDEN DE MÉRITO ACREDITADO", estilo="Sección anexo")
    tabla_fija(documento, [("PUESTO", "PERSONA / DNI")] + [(fila["puesto"], persona_texto(aspirantes[fila["aspirante_id"]]))
        for fila in datos["complementarios"]["orden_merito"]], [ancho // 6, ancho - ancho // 6])


def dictamen(documento, datos):
    apertura_resolucion(documento, datos, "APROBACIÓN DEL DICTAMEN")
    antecedentes_posteriores(documento, datos)
    resuelve(documento, datos)
    articulo(documento, 1, f'Aprobar el dictamen del jurado del concurso {modalidad(datos)} para el cargo de {cargo(datos)}, '
              f'convocado por resolución {datos["complementarios"]["res_convocatoria"]}, conforme al resultado y orden de mérito acreditados.')
    merito(documento, datos)
    cierre_resolucion(documento, datos, 2)


def desierto_inscriptos(documento, datos):
    """Borrador provisional por falta de inscriptos; no efectúa la nueva convocatoria."""
    c = datos["concurso"]
    siguiente = datos["nuevo_llamado"]["nombre"]
    apertura_resolucion(documento, datos, "DECLARACIÓN DE CONCURSO DESIERTO POR FALTA DE ASPIRANTES",
        marco_normativo="Que las actuaciones se sustanciaron conforme al Reglamento de Concursos "
        "del Personal Nodocente de la Universidad Nacional de San Juan, aprobado por "
        "Ordenanza N.º 33/10-CS y sus modificatorias.")
    parrafo(documento, f'Que mediante Resolución N.º {c["resolucion_convocatoria"]} se convocó '
        f'a Concurso {modalidad(datos)} de antecedentes y oposición para cubrir el cargo referido.')
    parrafo(documento, f'Que según el acta de cierre de inscripciones incorporada a fojas '
        f'{foja(datos, "acta_inscriptos")} del expediente de referencia, no se registraron aspirantes.')
    parrafo(documento, 'Que el artículo 45 del Anexo de la Ordenanza N.º 33/10-CS, '
        'texto actualizado por la Ordenanza N.º 21/13-CS, contempla la declaración de concurso '
        'desierto por falta de aspirantes y la realización de un nuevo llamado.')
    parrafo(documento, f'Que, de acuerdo con las modalidades previstas en el artículo 4 del mismo '
        f'Reglamento, corresponde tramitar el próximo llamado como Concurso {siguiente}.')
    # Separa los considerandos de la parte dispositiva del borrador.
    documento.add_page_break()
    resuelve(documento, datos)
    # Evitar dejar aislado «RESUELVE:» al pie de una página.
    for encabezado in documento.paragraphs[-3:]:
        encabezado.paragraph_format.keep_with_next = True
    articulo(documento, 1, f'Declarar DESIERTO, por falta de aspirantes, el Concurso {modalidad(datos)} '
        f'de antecedentes y oposición convocado mediante Resolución N.º {c["resolucion_convocatoria"]}, '
        f'para cubrir, con carácter efectivo, un cargo de Planta Permanente del Personal Nodocente '
        f'con funciones de “{c["denominacion"]}”, Categoría {c["categoria"]}, Tramo {c["tramo"]}, '
        f'Agrupamiento {agrupamiento_formateado(datos)}, Código Presupuestario '
        f'{c["codigo_presupuestario"]}, {estructura(datos)}, tramitado por Expediente '
        f'N.º {c["expediente"]}, conforme al acta de cierre de inscripciones obrante a '
        f'fojas {foja(datos, "acta_inscriptos")}.')
    articulo(documento, 2, f'Disponer la tramitación de un nuevo llamado a Concurso {siguiente}, '
        'de acuerdo con el Reglamento de Concursos del Personal Nodocente de la UNSJ, '
        'observando el plazo previsto en el artículo 45 de su Anexo. '
        'La presente disposición no sustituye el acto de convocatoria correspondiente.')
    articulo(documento, 3, 'Comuníquese a las dependencias intervinientes y déjese constancia '
        'en las actuaciones administrativas.')
    cierre_resolucion(documento, datos, 4)


def designacion(documento, datos):
    x = datos["complementarios"]
    apertura_resolucion(documento, datos, "DESIGNACIÓN")
    if vacancia := antecedente_vacancia(datos):
        parrafo(documento, vacancia)
    antecedentes_posteriores(documento, datos)
    parrafo(documento, "Que el dictamen fue aprobado por resolución oficial " + x["res_dictamen"] + ".")
    parrafo(documento, f'Que obra la afectación presupuestaria a fs. {foja(datos, "afectacion_presupuestaria")}.')
    merito(documento, datos)
    elegida = next(a for a in datos["aspirantes"] if str(a["id"]) == x["aspirante_id"])
    parrafo(documento, "Persona expresamente propuesta para designación: " + persona_texto(elegida["persona"]) + ".")
    resuelve(documento, datos)
    articulo(documento, 1, f'Designar a {persona_texto(elegida["persona"])}, en el cargo de {cargo(datos)}. {x["efectos"]}')
    numero = 2
    if x["aplica_limitacion"]:
        articulo(documento, numero, x["limitacion"])
        numero += 1
    articulo(documento, numero, x["circuito"])
    articulo(documento, numero + 1, "Imputar la erogación conforme al detalle de afectación presupuestaria acreditado.")
    presupuesto(documento, datos)
    cierre_resolucion(documento, datos, numero + 2)


def apertura_nota(documento, datos, destinatario):
    x = datos["complementarios"]
    p = documento.sections[0].header.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(str(Path(__file__).parent / "assets" / "unsj-crest-wordmark.png"), width=Cm(6.4))
    # Separación del membrete: no posiciones absolutas ni imágenes de firmas.
    p = parrafo(documento, x["lugar_fecha"], estilo="Metadato institucional")
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = parrafo(documento, destinatario, estilo="Metadato institucional")
    p.runs[0].bold = True
    parrafo(documento, "Presente", estilo="Metadato institucional")


def cierre_nota(documento, datos, texto):
    parrafo(documento, texto)
    p = parrafo(documento, datos["complementarios"]["firmante"] + "\nAutoridad firmante propuesta · Sin firma", estilo="Metadato institucional")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _cuerpo_oficio(documento, texto):
    p = documento.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(9)
    return p


def _encabezado_oficio(documento, datos, destinatarios, membrete):
    _membrete_notificacion(documento, datos["complementarios"]["lugar_fecha"], membrete)
    for texto in (*destinatarios, "Presente"):
        p = parrafo(documento, texto, estilo="Metadato institucional")
        p.paragraph_format.space_after = Pt(2)
    documento.paragraphs[-1].paragraph_format.space_after = Pt(20)


def _cargo_oficio(datos):
    c = datos["concurso"]
    return (f'Categoría {c["categoria"]}, Tramo {c["tramo"]}, Agrupamiento '
            f'{agrupamiento_formateado(datos)}, código presupuestario {c["codigo_presupuestario"]}, '
            f'para cumplir funciones de {c["denominacion"]}, {estructura(datos)}')


def _vacancia_oficio(datos):
    c = datos["concurso"]
    # Los detalles individuales del PDF original pertenecen a otro expediente:
    # nunca copiarlos ni presumir DNI/código de destino.
    frase = causa_de_vacancia(datos)
    if not frase:
        raise ValidationError("Completá los antecedentes de vacancia del concurso.")
    return frase[4:5].upper() + frase[5:] if frase.startswith("Que ") else frase


def elecciones(documento, datos, *, membrete=None):
    """Modelo de elección directa: párrafos fijos del PDF, hechos variables del concurso."""
    j = datos["destinatario_personal"]
    _encabezado_oficio(documento, datos, (
        f'{j["tratamiento"]} {j["titulo"]} {j["nombre"]} {j["apellido"]}'.replace("  ", " ").strip(),
        f'{j["cargo"]}',
    ), membrete)
    _cuerpo_oficio(documento,
        "Tengo el agrado de dirigirme a usted a fin de solicitar se instrumente la elección directa de "
        "los representantes del Personal de la dependencia (un miembro titular y un suplente) para integrar "
        f"el Jurado del Concurso {modalidad(datos)}, conforme a lo establecido en el Reglamento de "
        "Concursos del Personal Nodocente.")
    _cuerpo_oficio(documento,
        f'Dicho concurso se tramita por expediente Nº {datos["concurso"]["expediente"]} y tiene por objeto cubrir un '
        f'cargo vacante de planta permanente, {_cargo_oficio(datos)}.')
    _cuerpo_oficio(documento, _vacancia_oficio(datos))
    _cuerpo_oficio(documento,
        "Cabe señalar que, conforme a lo dispuesto por el Acta Paritaria N.º 2/13 —interpretativa del "
        "artículo 24 de la Ordenanza N.º 33/10-CS—, se establece el mecanismo de elección de los "
        "miembros del jurado en representación del Personal de la dependencia, así como la división de las "
        "unidades intervinientes en el ámbito del Rectorado.")
    _cuerpo_oficio(documento,
        "A tal efecto, el personal Nodocente, para ser elegible, deberá revestir carácter efectivo en un "
        "cargo de igual o superior categoría a la del cargo convocado.")
    _cuerpo_oficio(documento,
        "Se recuerda que el acto eleccionario deberá ser fiscalizado por la Asociación del Personal "
        "de la Universidad Nacional de San Juan (APUNSJ), conforme lo establecido en la normativa vigente.")
    _cuerpo_oficio(documento, "Sin otro motivo, saludo a usted atte.")
    _firma_oficio(documento)




def _firma_oficio(documento):
    p = documento.add_paragraph()
    p.paragraph_format.space_before = Pt(50)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Andrés A. Petronio\nDIRECTOR DE CONCURSOS NODOCENTES\nRECTORADO - UNSJ")


def _parrafo_notificacion(documento):
    """Párrafo justificado con la sangría y el espaciado del modelo institucional."""
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    return p


def _membrete_notificacion(documento, lugar_fecha, imagen):
    """Membrete institucional vigente subido por la autoridad, sin logos anuales fijos."""
    if not imagen:
        raise ValidationError("Antes de generar notificaciones, cargá el membrete vigente.")
    from docx.image.image import Image as ImagenDocx
    info = ImagenDocx.from_blob(imagen)
    relacion = info.px_width / info.px_height
    ancho = min(17.0, 3.1 * relacion)
    cabecera = documento.sections[0].header.paragraphs[0]
    cabecera.text = ""
    cabecera.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cabecera.paragraph_format.first_line_indent = Cm(0)
    cabecera.paragraph_format.space_after = Pt(0)
    cabecera.add_run().add_picture(BytesIO(imagen), width=Cm(ancho))
    fecha = parrafo(documento, lugar_fecha, estilo="Metadato institucional")
    fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fecha.paragraph_format.space_after = Pt(22)


def _representacion_para_nota(representacion):
    return {
        Jurado.Representacion.RECTORADO: "del Sr. Rector",
        Jurado.Representacion.DEPENDENCIA: "de la autoridad de la dependencia que convoca",
        Jurado.Representacion.PERSONAL: "del personal de la dependencia",
    }[representacion]


def notificacion_jurado(documento, datos, *, membrete=None):
    """Notificación individual conforme al modelo PDF de la Dirección de Concursos."""
    x, c, j = datos["complementarios"], datos["concurso"], datos["jurado_actual"]
    nom = j["persona"]
    destinatario = " ".join(v for v in (j["titulo"].strip(), nom["nombre"], nom["apellido"]) if v)
    # Sin atribuir género por el nombre ni pedir tratamientos nuevamente al operador.
    designacion = "designado/a"
    representacion = _representacion_para_nota(j["representacion"])
    _membrete_notificacion(documento, x["lugar_fecha"], membrete)
    # Encabezado del destinatario: título y nombre, unidad y «Presente».
    for texto in (destinatario, j["unidad"], "Presente"):
        p = parrafo(documento, texto, estilo="Metadato institucional")
        p.runs[0].bold = True
    documento.paragraphs[-1].paragraph_format.space_after = Pt(19)

    p = _parrafo_notificacion(documento)
    p.add_run("Tengo el agrado de dirigirme a usted a fin de notificarle que, mediante Resolución Nº ")
    p.add_run(x["res_convocatoria"])
    p.add_run(", ha sido ")
    p.add_run(designacion)
    p.add_run(" en carácter de ")
    p.add_run(f'{j["calidad"]}, en representación {representacion}').bold = True
    p.add_run(", para integrar el jurado del Concurso ")
    p.add_run(modalidad(datos))
    p.add_run(" destinado a cubrir, con carácter efectivo y mediante evaluación de antecedentes y oposición, "
              "un cargo de planta permanente del personal Nodocente, Categoría ")
    p.add_run(str(c["categoria"]))
    p.add_run(", Tramo ")
    p.add_run(c["tramo"])
    p.add_run(" del Agrupamiento ")
    p.add_run(agrupamiento_formateado(datos))
    p.add_run(", Código Presupuestario Nº ")
    p.add_run(c["codigo_presupuestario"])
    p.add_run(", con funciones de “")
    p.add_run(c["denominacion"])
    p.add_run("” ")
    p.add_run(estructura(datos))
    p.add_run(".")

    p = _parrafo_notificacion(documento)
    p.add_run("El mencionado concurso se sustancia conforme a los términos y condiciones establecidos "
              "en las Ordenanzas N.º 33/10-CS y 20/11-CS, y tramita por Expediente N.º ")
    p.add_run(c["expediente"])
    p.add_run(".")

    p = _parrafo_notificacion(documento)
    p.add_run("Asimismo, se solicita que tenga a bien informar una dirección de correo electrónico a los "
              "efectos de remitirle copia de la Resolución Nº ")
    p.add_run(x["res_convocatoria"])
    p.add_run(" y cursar las comunicaciones y notificaciones correspondientes al desarrollo del concurso.")
    p = _parrafo_notificacion(documento)
    p.add_run("Sin otro particular, saludo a usted atentamente.")

    # Área libre para sello y firma ológrafa: no imprime un cargo institucional prefijado.
    espacio_firma = documento.add_paragraph(style="Metadato institucional")
    espacio_firma.paragraph_format.space_before = Pt(12)
    espacio_firma.paragraph_format.space_after = Pt(38)
    espacio_firma.paragraph_format.line_spacing = 1
    for etiqueta, linea in (("Firma", "..............................................."),
                           ("Aclaración", "______________________________________"),
                           ("Fecha", "____/____/________"),
                           ("Correo electrónico", "______________________________________")):
        p = parrafo(documento, f"{etiqueta}: {linea}", estilo="Metadato institucional")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_together = True



RENDERERS.update(convocatoria=convocatoria, jurados=notificacion_jurado, dictamen=dictamen,
                 desierto_inscriptos=desierto_inscriptos,
                 designacion=designacion, elecciones=elecciones)
